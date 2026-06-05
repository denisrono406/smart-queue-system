from datetime import datetime, timezone, timedelta
from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file
from flask_login import login_required, current_user
from app.extensions import db
from app.models import (User, Role, Organization, ServiceDesk, Ticket, Report)
from app.forms import ReportForm
from app.utils import calculate_estimated_wait
import io

reports_bp = Blueprint('reports', __name__)


def admin_required(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.can_manage_org():
            flash('Access denied.', 'danger')
            return redirect(url_for('main.index'))
        return f(*args, **kwargs)
    return decorated


@reports_bp.route('/')
@login_required
@admin_required
def index():
    org = current_user.organization
    reports = Report.query.filter_by(organization_id=org.id).order_by(
        Report.created_at.desc()
    ).limit(20).all()
    return render_template('reports/index.html', reports=reports)


@reports_bp.route('/generate', methods=['GET', 'POST'])
@login_required
@admin_required
def generate():
    org = current_user.organization
    form = ReportForm()

    desks = ServiceDesk.query.filter_by(organization_id=org.id).all()
    form.service_desk_id.choices = [('', '-- All Desks --')] + [
        (d.id, d.name) for d in desks
    ]

    if form.validate_on_submit():
        report_type = form.report_type.data
        date_from = form.date_from.data
        date_to = form.date_to.data
        file_format = form.file_format.data
        desk_id = form.service_desk_id.data or None

        # Query tickets for the report period
        query = Ticket.query.filter(
            Ticket.organization_id == org.id,
            Ticket.joined_at >= date_from,
            Ticket.joined_at <= date_to + timedelta(days=1)
        )
        if desk_id:
            query = query.filter_by(service_desk_id=desk_id)

        tickets = query.all()

        # Calculate summary
        total = len(tickets)
        completed = len([t for t in tickets if t.status == 'completed'])
        cancelled = len([t for t in tickets if t.status == 'cancelled'])
        no_show = len([t for t in tickets if t.status == 'no_show'])
        waiting = len([t for t in tickets if t.status == 'waiting'])

        avg_wait = 0
        completed_with_times = [t for t in tickets if t.status == 'completed' and t.served_at and t.joined_at]
        if completed_with_times:
            total_wait = sum((t.served_at - t.joined_at).total_seconds() / 60 for t in completed_with_times)
            avg_wait = round(total_wait / len(completed_with_times), 1)

        summary = {
            'total_tickets': total,
            'completed': completed,
            'cancelled': cancelled,
            'no_show': no_show,
            'waiting': waiting,
            'avg_wait_minutes': avg_wait,
            'completion_rate': round(completed / total * 100, 1) if total > 0 else 0,
        }

        # Desk performance breakdown
        desk_performance = {}
        for ticket in tickets:
            desk_name = ticket.service_desk.name if ticket.service_desk else 'Unknown'
            if desk_name not in desk_performance:
                desk_performance[desk_name] = {'total': 0, 'completed': 0, 'avg_wait': 0, 'wait_times': []}
            desk_performance[desk_name]['total'] += 1
            if ticket.status == 'completed':
                desk_performance[desk_name]['completed'] += 1
            if ticket.served_at and ticket.joined_at:
                desk_performance[desk_name]['wait_times'].append(
                    (ticket.served_at - ticket.joined_at).total_seconds() / 60
                )

        for name, data in desk_performance.items():
            if data['wait_times']:
                data['avg_wait'] = round(sum(data['wait_times']) / len(data['wait_times']), 1)

        # Generate file
        if file_format == 'pdf':
            output = _generate_pdf_report(org, report_type, date_from, date_to,
                                          summary, desk_performance, tickets)
            filename = f'report_{report_type}_{date_from}_{date_to}.pdf'
            mimetype = 'application/pdf'
        else:
            output = _generate_excel_report(org, report_type, date_from, date_to,
                                             tickets, summary, desk_performance)
            filename = f'report_{report_type}_{date_from}_{date_to}.xlsx'
            mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

        # Save report record
        report = Report(
            organization_id=org.id,
            generated_by=current_user.id,
            report_type=report_type,
            date_from=date_from,
            date_to=date_to,
            file_format=file_format,
            summary=summary,
        )
        db.session.add(report)
        db.session.commit()

        return send_file(output, as_attachment=True,
                         download_name=filename, mimetype=mimetype)

    return render_template('reports/generate.html', form=form)


def _generate_pdf_report(org, report_type, date_from, date_to,
                          summary, desk_performance, tickets):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph(f"<b>{org.name}</b>", styles['Title']))
    story.append(Paragraph(f"{report_type.title()} Report: {date_from} to {date_to}", styles['Heading2']))
    story.append(Spacer(1, 0.3*inch))

    # Summary table
    summary_data = [
        ['Metric', 'Value'],
        ['Total Tickets', str(summary['total_tickets'])],
        ['Completed', str(summary['completed'])],
        ['Cancelled', str(summary['cancelled'])],
        ['No Show', str(summary['no_show'])],
        ['Average Wait (min)', str(summary['avg_wait_minutes'])],
        ['Completion Rate', f"{summary['completion_rate']}%"],
    ]
    table = Table(summary_data, colWidths=[3*inch, 2*inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a365d')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(table)
    story.append(Spacer(1, 0.3*inch))

    # Desk performance table
    if desk_performance:
        story.append(Paragraph("<b>Desk Performance</b>", styles['Heading3']))
        desk_data = [['Desk', 'Total', 'Completed', 'Avg Wait (min)']]
        for name, data in desk_performance.items():
            desk_data.append([name, str(data['total']), str(data['completed']), str(data['avg_wait'])])
        desk_table = Table(desk_data, colWidths=[2.5*inch, 1*inch, 1*inch, 1.5*inch])
        desk_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2d3748')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('PADDING', (0, 0), (-1, -1), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        story.append(desk_table)

    doc.build(story)
    buffer.seek(0)
    return buffer


def _generate_excel_report(org, report_type, date_from, date_to,
                            tickets, summary, desk_performance):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    buffer = io.BytesIO()
    wb = Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = 'Summary'
    header_fill = PatternFill(start_color='1A365D', end_color='1A365D', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    ws.append(['Metric', 'Value'])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for key, val in summary.items():
        ws.append([key.replace('_', ' ').title(), val])

    # Tickets detail sheet
    ws2 = wb.create_sheet('Tickets')
    headers = ['Ticket #', 'User', 'Desk', 'Priority', 'Status', 'Position',
               'Est. Wait (min)', 'Joined At', 'Completed At']
    ws2.append(headers)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font

    for t in tickets:
        ws2.append([
            t.ticket_number,
            t.user.full_name if t.user else '',
            t.service_desk.name if t.service_desk else '',
            t.priority_label,
            t.status,
            t.queue_position,
            t.estimated_wait_minutes,
            t.joined_at.strftime('%Y-%m-%d %H:%M') if t.joined_at else '',
            t.completed_at.strftime('%Y-%m-%d %H:%M') if t.completed_at else '',
        ])

    # Desk performance sheet
    if desk_performance:
        ws3 = wb.create_sheet('Desk Performance')
        ws3.append(['Desk', 'Total', 'Completed', 'Avg Wait (min)'])
        for cell in ws3[1]:
            cell.fill = header_fill
            cell.font = header_font

        for name, data in desk_performance.items():
            ws3.append([name, data['total'], data['completed'], data['avg_wait']])

    wb.save(buffer)
    buffer.seek(0)
    return buffer
